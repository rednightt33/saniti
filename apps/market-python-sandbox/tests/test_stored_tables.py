"""D0 (round 2026-10-03): a stored table paged, exported or recounted without a session, from the file market-ai-orc
keeps for the life of the conversation."""
from __future__ import annotations

import base64
import hashlib
import io
import json

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from app import stored_tables
from app.stored_tables import StoredTableError
from conftest import requires_root
from test_conversation_reuse import reuse  # noqa: F401
from test_dataneed_bundles import HEADERS


def parquet(rows: list[dict]) -> bytes:
    out = io.BytesIO()
    pq.write_table(pa.Table.from_pylist(rows), out)
    return out.getvalue()


def meta_for(data: bytes, **extra) -> dict:
    return {"format": "PARQUET", "checksum_sha256": hashlib.sha256(data).hexdigest(), **extra}


# the g5 shape: one row per crash day for broker RB, net value of that day
CRASH = [{"date": f"2026-0{1 + i // 28}-{1 + i % 28:02d}", "broker": "RB", "net_value": (i % 9) - 2.0}
         for i in range(132)]


def test_page_reads_rows_like_a_session_output() -> None:
    data = parquet(CRASH)
    got = stored_tables.page(data, meta_for(data), 130, 50)
    assert got["row_count"] == 132 and len(got["rows"]) == 2 and got["next_offset"] is None
    first = stored_tables.page(data, meta_for(data), 0, 100)
    assert first["next_offset"] == 100 and first["rows"][0]["broker"] == "RB"


def test_a_wrong_checksum_or_format_is_refused() -> None:
    data = parquet(CRASH)
    with pytest.raises(StoredTableError) as bad:
        stored_tables.page(data, {**meta_for(data), "checksum_sha256": "0" * 64}, 0, 10)
    assert bad.value.code == "STORED_TABLE_CHECKSUM_MISMATCH"
    with pytest.raises(StoredTableError) as fmt:
        stored_tables.page(data, {**meta_for(data), "format": "PNG"}, 0, 10)
    assert fmt.value.code == "STORED_TABLE_FORMAT"


def test_recount_reproduces_a_count_claim_with_its_rows() -> None:
    """The g5 claim "RB net buy on N of 132 crash days" recomputed from the base table."""
    data = parquet(CRASH)
    expected = sum(1 for r in CRASH if r["net_value"] > 0)
    got = stored_tables.recount(data, meta_for(data, where=[{"column": "net_value", "op": "GT", "value": 0}],
                                               measure="COUNT"))
    assert got["value"] == expected and got["rows_matched"] == expected and got["rows_total"] == 132
    assert all(r["net_value"] > 0 for r in got["rows"])


def test_recount_other_measures_and_conditions() -> None:
    """Cases other than the observed one: a price low over a date range, distinct brokers, a sum by text value."""
    prices = [{"ticker": t, "date": f"2026-09-{d:02d}", "low": 100.0 + d + (5 if t == "BMRI" else 0)}
              for t in ("BBCA", "BMRI") for d in range(1, 21)]
    data = parquet(prices)
    low = stored_tables.recount(data, meta_for(data, measure="MIN", column="low", where=[
        {"column": "ticker", "op": "EQ", "value": "BMRI"}, {"column": "date", "op": "GE", "value": "2026-09-10"}]))
    assert low["value"] == 115.0
    distinct = stored_tables.recount(data, meta_for(data, measure="COUNT_DISTINCT", column="ticker"))
    assert distinct["value"] == 2
    both = stored_tables.recount(data, meta_for(data, measure="SUM", column="low", where=[
        {"column": "ticker", "op": "IN", "values": ["BBCA"]}]))
    assert both["value"] == sum(101.0 + d - 1 for d in range(1, 21))


def test_recount_refuses_an_unknown_column_with_the_columns() -> None:
    data = parquet(CRASH)
    with pytest.raises(StoredTableError) as bad:
        stored_tables.recount(data, meta_for(data, measure="SUM", column="net"))
    assert bad.value.code == "RECOUNT_INVALID" and "net_value" in bad.value.details["columns"]


@pytest.mark.parametrize("target", ["CSV", "PARQUET", "XLSX"])
def test_export_keeps_every_row_and_adds_definition_and_lineage(target: str) -> None:
    data = parquet(CRASH)
    body, mime, ext = stored_tables.export(data, meta_for(
        data, target=target, definition={"filters": ["Broker = RB"], "period": "crash days"},
        lineage={"source_tables": ["IDX_Broker_Summary"], "data_as_of": "2026-09-30"}))
    assert ext == target.lower() and mime
    if target == "CSV":
        assert body.decode().count("\n") == 133
    elif target == "PARQUET":
        assert pq.read_table(io.BytesIO(body)).num_rows == 132
    else:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(body), read_only=True)
        assert book.sheetnames == ["data", "definisi", "asal data"]
        assert sum(1 for _ in book["data"].iter_rows()) == 133
        lineage = {row[0]: row[1] for row in book["asal data"].iter_rows(values_only=True)}
        assert lineage["source_tables"] == "IDX_Broker_Summary"


def test_an_xlsx_shows_the_readers_column_titles_and_a_column_sheet() -> None:
    """M128b (2026-10-09): the Excel export of BBRI's foreign flow showed foreign_net_value and turnover_idr as
    headers; the caller's titles and column meanings make the file readable, the CSV keeps the machine names."""
    from openpyxl import load_workbook

    data = parquet(CRASH)
    names = pq.read_table(io.BytesIO(data)).column_names
    headers = {names[0]: "Tanggal", names[-1]: "Nilai bersih (Rp)"}
    columns = [{"label": "Tanggal", "name": names[0], "meaning": "Hari bursa", "unit": None},
               {"label": "Nilai bersih (Rp)", "name": names[-1], "meaning": "Beli dikurangi jual", "unit": "IDR"}]
    body, _, _ = stored_tables.export(data, meta_for(data, target="XLSX", headers=headers, columns=columns))
    book = load_workbook(io.BytesIO(body), read_only=True)
    assert book.sheetnames == ["data", "kolom"]
    head = next(book["data"].iter_rows(values_only=True))
    assert head[0] == "Tanggal" and head[-1] == "Nilai bersih (Rp)" and list(head[1:-1]) == names[1:-1]
    assert list(book["kolom"].iter_rows(values_only=True))[2] == ("Nilai bersih (Rp)", names[-1], "Beli dikurangi jual",
                                                                  "IDR")
    csv, _, _ = stored_tables.export(data, meta_for(data, target="CSV", headers=headers, columns=columns))
    assert csv.decode().splitlines()[0].replace('"', "").split(",") == names


TEXTS = {"row_rule_label": "Aturan baris", "row_rule_one": "Satu baris untuk setiap {key}.",
         "row_rule_many": "Satu baris untuk setiap {keys}.", "and": " dan ",
         "row_rule_none": "Tidak ada kolom yang unik per baris.", "completeness_label": "Kelengkapan",
         "period": "{start} s.d. {end}, {calendar} hari bursa.", "group": "{group}: {file} dari {source} hari di sumber.",
         "all_rows": "Semua baris", "unchecked": "Tidak dapat dicek otomatis.",
         "activity_only": "Sumber hanya mencatat hari dengan aktivitas."}


def definition_rows(body: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    return [r for r in load_workbook(io.BytesIO(body), read_only=True)["definisi"].iter_rows(values_only=True)][1:]


def test_an_xlsx_says_what_one_row_is_and_how_complete_each_group_is() -> None:
    """M128 (d): BBRI foreign flow per board; Nego had a trading day without a foreign trade, so the file holds one
    day fewer than the source. The definition sheet says one row per date and board, and Nego 2 of 3 days."""
    rows = [{"date": d, "board": b, "net": 1.0} for d in ("2026-05-25", "2026-05-26", "2026-05-27")
            for b in ("Regular", "Nego") if not (b == "Nego" and d == "2026-05-26")]
    data = parquet(rows)
    check = {"status": "CHECKED", "time_column": "date", "group_columns": ["board"], "from": "2026-05-25",
             "to": "2026-05-27", "calendar_dates": 3, "row_presence": "ACTIVITY_ONLY",
             "groups": [{"key": ["Regular"], "days": 3}, {"key": ["Nego"], "days": 3}]}
    body, _, _ = stored_tables.export(data, meta_for(data, target="XLSX", texts=TEXTS, completeness=check,
                                                     headers={"date": "Tanggal", "board": "Papan"}))
    got = definition_rows(body)
    assert ("Aturan baris", "Satu baris untuk setiap Tanggal dan Papan.") in got
    assert ("Kelengkapan", "Papan Regular: 3 dari 3 hari di sumber.") in got
    assert ("Kelengkapan", "Papan Nego: 2 dari 3 hari di sumber.") in got
    assert ("Kelengkapan", "Sumber hanya mencatat hari dengan aktivitas.") in got


def test_completeness_of_other_shapes_and_when_it_cannot_be_checked() -> None:
    """Cases other than the observed one: one ticker's daily prices without a group (dates as date objects), a file
    whose rows repeat, and a check the caller could not run."""
    from datetime import date as day

    prices = [{"date": day(2026, 9, d), "close": 100.0 + d} for d in range(1, 6)]
    data = parquet(prices)
    check = {"status": "CHECKED", "time_column": "date", "group_columns": [], "from": "2026-09-02",
             "to": "2026-09-05", "calendar_dates": 4, "groups": [{"key": [], "days": 4}]}
    got = definition_rows(stored_tables.export(data, meta_for(data, target="XLSX", texts=TEXTS,
                                                              completeness=check))[0])
    assert ("Aturan baris", "Satu baris untuk setiap date.") in got
    assert ("Kelengkapan", "Semua baris: 4 dari 4 hari di sumber.") in got
    repeated = parquet([{"broker": "RB", "net": 1.0}, {"broker": "RB", "net": 2.0}])
    got = definition_rows(stored_tables.export(repeated, meta_for(repeated, target="XLSX", texts=TEXTS,
                                                                  completeness={"status": "UNCHECKED"}))[0])
    assert got == [("Aturan baris", "Tidak ada kolom yang unik per baris."),
                   ("Kelengkapan", "Tidak dapat dicek otomatis.")]


def test_recount_reads_the_first_and_last_date_of_a_date_column() -> None:
    from datetime import date as day

    for rows in (CRASH, [{"date": day(2026, 1, d), "v": 1.0} for d in (3, 1, 2)]):
        data = parquet(rows)
        first = stored_tables.recount(data, meta_for(data, measure="MIN", column="date"))["value"]
        last = stored_tables.recount(data, meta_for(data, measure="MAX", column="date"))["value"]
        assert (first, last) in (("2026-01-01", "2026-05-20"), ("2026-01-01", "2026-01-03"))
    assert stored_tables.recount(data, meta_for(data, measure="SUM", column="date"))["value"] is None


def test_export_over_the_limit_is_refused_with_a_next_step() -> None:
    data = parquet(CRASH)
    with pytest.raises(StoredTableError) as big:
        stored_tables.export(data, meta_for(data, target="CSV", max_bytes=100))
    assert big.value.code == "EXPORT_TOO_LARGE" and big.value.http_status == 413 and big.value.next_action


def test_json_rows_are_a_table_for_export() -> None:
    """Evidence rows (JSON) can be exported like an output table."""
    data = json.dumps({"rows": CRASH[:3]}).encode()
    body, _, _ = stored_tables.export(data, {"format": "JSON", "checksum_sha256": hashlib.sha256(data).hexdigest(),
                                             "target": "CSV"})
    assert body.decode().splitlines()[0].replace('"', "") == "date,broker,net_value"


@requires_root
def test_the_routes_and_the_runtime_capability(reuse) -> None:  # noqa: F811
    api = reuse["api"]
    runtime = api.get("/v1/runtime", headers=HEADERS).json()["stored_tables"]
    assert runtime["enabled"] is True and runtime["ops"] == ["page", "export", "recount"]
    data = parquet(CRASH)

    def call(op: str, meta: dict):
        encoded = base64.urlsafe_b64encode(json.dumps(meta).encode()).decode().rstrip("=")
        return api.post(f"/v1/stored-tables/{op}", content=data, headers={
            **HEADERS, "X-Saniti-Output-Meta": encoded, "Content-Type": "application/octet-stream"})

    assert call("page", meta_for(data, offset=0, limit=5)).json()["rows"][0]["broker"] == "RB"
    exported = call("export", meta_for(data, target="XLSX"))
    assert exported.status_code == 200
    assert hashlib.sha256(exported.content).hexdigest() == exported.headers["X-Saniti-Checksum-Sha256"]
    assert exported.headers["X-Saniti-Extension"] == "xlsx"
    refused = call("recount", meta_for(data, measure="MEAN", column="missing"))
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "RECOUNT_INVALID"
    assert call("nope", meta_for(data)).status_code == 404
