"""E9 (EXEC-1018B, M176; user decisions 2026-10-16 "pakai E dengan masa penyimpanan 1 bulan ... lewat 1 bulan ... go for
option F"): the data behind a result can be handed to the user. A completed analysis releases its working tables
(save_table) and the tables marked emit_table(..., support=True) as supporting tables, byte for byte; a bundle's
delivered rows of one data request come back as one file; an approved need is carried into a later request of the same
conversation without the model, so the same rows can be extracted again."""
from __future__ import annotations

import base64
import io
import json
import shutil

import pyarrow.parquet as pq

from conftest import requires_root
from test_conversation_reuse import KEY, OTHER, first_turn, headers, reuse  # noqa: F401 - reuse is a fixture
from test_dataneed_bundles import HEADERS, env  # noqa: F401 - env is a fixture
from test_dataneed_sessions import ok, session  # noqa: F401 - session is a fixture

pytestmark = requires_root


def complete(env):  # noqa: F811
    return env["api"].post(f"/v1/sessions/{env['session_id']}/complete", json={"request_id": "req_bundle_1"},
                           headers=HEADERS).json()


def test_working_tables_and_marked_tables_are_released_as_supporting_tables(session) -> None:  # noqa: F811
    ok(session, "prices = load('prices')\nbanks = load('stock_classification')\n"
                "save_table('daily_rows', prices)\n"
                "emit_table('step_one', prices.head(3), support=True)\n"
                "emit_table('last_close', prices.groupby('ticker', as_index=False)['close'].last(), definition={})")
    result = complete(session)
    assert result["status"] == "COMPLETED", result  # a supporting table needs no definition
    support = {o["name"]: o for o in result["released_outputs"] if o.get("support")}
    assert set(support) == {"daily_rows", "step_one"}
    assert not [o for o in result["released_outputs"] if o["name"] == "last_close" and o.get("support")]
    import hashlib

    stored = session["dataneed"].store.get_output(support["daily_rows"]["output_id"])
    data = (session["dataneed"].sessions.outputs_root / stored["relative_path"]).read_bytes()
    assert support["daily_rows"]["checksum_sha256"] == hashlib.sha256(data).hexdigest()
    assert pq.read_table(io.BytesIO(data)).num_rows == support["daily_rows"]["row_count"] > 3
    assert support["daily_rows"]["lineage"]["need_id"]


def test_a_bundle_dataset_comes_back_as_one_file_of_its_approved_ranges(reuse) -> None:  # noqa: F811
    turn = first_turn(reuse)
    path = f"/v1/bundles/{turn['bundle_id']}/datasets/data_request_1_A/file"
    answer = reuse["api"].get(path, params={"request_id": "req_turn_2"}, headers=headers())
    assert answer.status_code == 200, answer.text
    table = pq.read_table(io.BytesIO(answer.content))
    meta = json.loads(base64.urlsafe_b64decode(answer.headers["x-saniti-dataset-meta"] + "=="))
    assert table.num_rows == meta["rows"] > 0 and meta["need_id"] == turn["need"]["need_id"]
    time_column = next(d["time_column"] for d in reuse["dataneed"].store.get_bundle(turn["bundle_id"])["manifest"]
                       ["datasets"] if d["data_request_id"] == "data_request_1_A")
    dates = sorted({str(d) for d in table.column(time_column).to_pylist()})
    starts = min(r["start"] for r in meta["ranges"])
    assert not dates or dates[0] >= starts  # warm-up buffer rows are left out
    other = reuse["api"].get(path, params={"request_id": "req_turn_2"}, headers=headers(OTHER))
    assert other.status_code == 404
    unknown = reuse["api"].get(f"/v1/bundles/{turn['bundle_id']}/datasets/nope/file",
                               params={"request_id": "req_turn_2"}, headers=headers())
    assert unknown.status_code == 404 and unknown.json()["error"]["code"] == "DATA_REQUEST_NOT_FOUND"
    shutil.rmtree(reuse["dataneed"].bundles.root / turn["bundle_id"])
    gone = reuse["api"].get(path, params={"request_id": "req_turn_2"}, headers=headers())
    assert gone.status_code == 410 and gone.json()["error"]["code"] == "BUNDLE_EXPIRED"
    assert gone.json()["error"]["rows_identity"] == meta["rows_identity"]  # the rebuilt rows are compared to this


def test_an_approved_need_is_carried_into_a_later_request_of_the_same_conversation(reuse) -> None:  # noqa: F811
    turn = first_turn(reuse)
    need_id = turn["need"]["need_id"]
    carried = reuse["api"].post(f"/v1/data-needs/{need_id}/carry", json={"request_id": "req_turn_9"},
                                headers=headers()).json()
    assert carried["status"] == "APPROVED" and carried["need_id"] != need_id and carried["carried_from"] == need_id
    again = reuse["dataneed"].get_need(carried["need_id"])
    assert again["request_id"] == "req_turn_9" and again["spec_sha256"] == turn["need"]["spec_sha256"]
    assert [r["scope_sha256"] for r in again["requests"].values()] == \
        [r["scope_sha256"] for r in turn["need"]["requests"].values()]
    refused = reuse["api"].post(f"/v1/data-needs/{need_id}/carry", json={"request_id": "req_turn_9"},
                                headers=headers(OTHER))
    assert refused.status_code == 404 and refused.json()["error"]["code"] == "NEED_NOT_FOUND"


def _export(reuse, table_bytes: bytes, inputs: list[dict]):  # noqa: F811
    import hashlib

    meta = {"format": "PARQUET", "checksum_sha256": hashlib.sha256(table_bytes).hexdigest(), "target": "XLSX",
            "request_id": "req_turn_2", "inputs": inputs, "definition": {"notes": "x"},
            "texts": {"inputs": "Input rows in {sheets}.", "inputs_field": "Lembar input"}}
    encoded = base64.urlsafe_b64encode(json.dumps(meta).encode()).decode().rstrip("=")
    return reuse["api"].post("/v1/stored-tables/export", content=table_bytes,
                             headers={**headers(), "X-Saniti-Output-Meta": encoded,
                                      "Content-Type": "application/octet-stream"})


def test_an_xlsx_gets_a_sheet_for_each_input_its_rows_do_not_cover(reuse) -> None:  # noqa: F811
    """E1: a result without the input's dates (a summary) gets the input rows, warm-up included, as a sheet; a result
    that spans the input's dates gets none; an input whose bundle is gone is reported, not guessed."""
    import pyarrow as pa
    from openpyxl import load_workbook

    turn = first_turn(reuse)
    item = {"bundle_id": turn["bundle_id"], "data_request_id": "data_request_1_A", "sheet": "prices"}
    summary = io.BytesIO()
    pq.write_table(pa.table({"ticker": ["BBRI"], "ret": [0.1]}), summary)
    answer = _export(reuse, summary.getvalue(), [item])
    assert answer.status_code == 200, answer.text
    statuses = json.loads(base64.urlsafe_b64decode(answer.headers["x-saniti-inputs"] + "=="))
    assert statuses[0]["status"] == "INCLUDED" and statuses[0]["rows"] > 0
    book = load_workbook(io.BytesIO(answer.content), read_only=True)
    assert "input prices" in book.sheetnames
    assert any("input prices" in str(row) for row in book["definisi"].iter_rows(values_only=True))
    full = reuse["api"].get(f"/v1/bundles/{turn['bundle_id']}/datasets/data_request_1_A/file",
                            params={"request_id": "req_turn_2"}, headers=headers())
    whole = reuse["dataneed"].dataset_file(turn["bundle_id"], "data_request_1_A", "req_turn_2", KEY, buffers=True)[0]
    assert pq.read_table(io.BytesIO(whole)).num_rows >= pq.read_table(io.BytesIO(full.content)).num_rows
    covered = _export(reuse, whole, [item])
    assert json.loads(base64.urlsafe_b64decode(covered.headers["x-saniti-inputs"] + "=="))[0]["status"] == "COVERED"
    shutil.rmtree(reuse["dataneed"].bundles.root / turn["bundle_id"])
    gone = _export(reuse, summary.getvalue(), [item])
    assert json.loads(base64.urlsafe_b64decode(gone.headers["x-saniti-inputs"] + "=="))[0] == {
        "data_request_id": "data_request_1_A", "sheet": "prices", "status": "NOT_INCLUDED",
        "reason": "BUNDLE_EXPIRED"}


def test_an_unchanged_working_table_is_not_copied_again_and_a_large_one_is_named(session, monkeypatch) -> None:  # noqa: F811,E501
    from app import sessions as session_module

    ok(session, "p = load('prices')\nsave_table('kept', p.head(5))\nsave_table('big', p)\n"
                "load('stock_classification')\nemit_json('n', {'n': 1}, definition={})")
    small = session["dataneed"].sessions.settings.jobs_dir
    from pathlib import Path

    big = (Path(small) / session["session_id"] / "intermediate" / "tables" / "big.parquet").stat().st_size
    monkeypatch.setattr(session_module, "SUPPORT_TABLE_BYTES", big - 1)
    first = complete(session)
    assert {o["name"] for o in first["released_outputs"] if o.get("support")} == {"kept"}
    assert first["supporting_tables_not_kept"] == ["big"]


def test_a_working_table_released_unchanged_is_not_copied_at_a_later_completion(reuse) -> None:  # noqa: F811
    from test_conversation_reuse import approve, complete, execute, post
    from test_dataneed_bundles import build, ytd_parts

    need = approve(reuse, "req_turn_1")
    bundle = build(reuse, need, ytd_parts(reuse, need), request_id="req_turn_1").json()
    sid = post(reuse, "/v1/sessions", {"request_id": "req_turn_1", "bundle_id": bundle["input_bundle_id"]}).json()[
        "session_id"]
    execute(reuse, sid, "req_turn_1", "p = load('prices')\nload('stock_classification')\n"
                                      "save_table('kept', p.head(5))\nemit_json('n', {'n': 1}, definition={})")
    first = complete(reuse, sid, "req_turn_1")
    assert [o["name"] for o in first["released_outputs"] if o.get("support")] == ["kept"]
    execute(reuse, sid, "req_turn_1", "save_table('later', p.head(2))\nemit_json('m', {'m': 2}, definition={})")
    again = complete(reuse, sid, "req_turn_1")
    assert again["status"] == "COMPLETED", again
    assert [o["name"] for o in again["released_outputs"] if o.get("support")] == ["later"]
