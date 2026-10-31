#!/usr/bin/env python3
"""The final Flow indicator workbook for the user (request 2026-10-31: "Excel finalnya").

FLOW_INDICATORS_SPEC.xlsx is the agreed contract (user decisions 2026-10-30) and stays unchanged: the Terminal's registry
is generated from it (`scripts/build_flow_registry.py --check`). This script copies its sheets as they are and adds:
- Rumus_Kode: the exact formula of every indicator and the shared conventions, read from FORMULAS.md (what
  `apps/edge-bff/terminal/src/calc/flow.ts` computes);
- Perubahan: what changed after the spreadsheet was agreed (reviews and decisions of 2026-10-30 and 2026-10-31);
- Versi: the sources, their sha256 and the registry check.

  python3 verification/terminal_api/flow_indicators/build_final_workbook.py <out.xlsx>
"""
import hashlib
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SPEC = HERE / 'FLOW_INDICATORS_SPEC.xlsx'
FORMULAS = HERE / 'FORMULAS.md'
REGISTRY = ROOT / 'apps/edge-bff/terminal/src/data/flowRegistry.json'
FLOW_TS = ROOT / 'apps/edge-bff/terminal/src/calc/flow.ts'

CHANGES = [
    ('2026-10-30', 'Semua (pemanggil Screener/Backtest)', 'Parameter yang tidak dikirim memakai default di registry, '
     'bukan nilai cadangan di dalam rumus.', 'Reviewer A, A4: FL26 tanpa `scope` terbaca "all" di satu tempat dan '
     '"topn" di tempat lain.', 'Terminal tidak berubah (selalu mengirim semua parameter).'),
    ('2026-10-30', 'FL09', 'Ambang "> 0" tidak menerima 0; dialog menulis "Above 0".', 'Reviewer A, A5.',
     'Nilai 0 ditolak di dialog.'),
    ('2026-10-30', 'FL18-FL20', 'Dicatat: Top-N beku diurutkan atas seluruh periode di chart, jadi melihat candle sesudah t.',
     'Reviewer A, A8.', 'Rumus sama; Backtest dan Screener tidak boleh membacanya sebagai informasi yang diketahui di t.'),
    ('2026-10-30', 'FL19, FL20', 'Aturan inventori di 0 atau di bawahnya mengikuti opsi A.', 'Keputusan user: "1 pakai A".',
     'Aturan live tetap; tidak ada perubahan kode.'),
    ('2026-10-30', 'FL26', 'Dicatat: candle tanpa data = kosong; candle dengan data tetapi tanpa net = 0.', 'Reviewer A, A7.',
     'Rumus sama.'),
    ('2026-10-30', 'Semua', 'Sesi broker pada tanggal yang tidak punya harga tidak masuk candle mana pun.',
     'Reviewer A, A1 (celah 17-18 Sep, M238).', 'Celah 17-18 Sep 2026 diisi 2026-10-31 (EXEC-PRICE-GAP-0917); '
     'Flow kini menghitung broker kedua hari itu. Aturannya tetap untuk celah lain.'),
    ('2026-10-30', 'FL08, FL21, FL37', 'Garis referensi tanpa default tetap punya isian di dialog.', 'Reviewer B, B1.',
     'Tampilan saja.'),
    ('2026-10-30', 'Semua (W/M)', 'Di chart mingguan/bulanan label "(1 = daily)" dan "Daily" mengikuti candle '
     '("Rolling weeks (1 = weekly)").', 'Reviewer B, B2-B3.', 'Label saja; rumus per candle sudah begitu.'),
    ('2026-10-30', 'FL09', 'Label ambang partisipasi: "Buy participation (%) at least".', 'Reviewer B, B4.', 'Label saja.'),
    ('2026-10-30', 'FL14', 'Empat desimal seperti FL13.', 'Reviewer B, B5.', 'Tampilan saja.'),
    ('2026-10-30', 'FL01-FL23, FL34', 'Tooltip dalam bahasa Indonesia sederhana; FL34 memberi tahu bila klasifikasi '
     'belum dimuat; dialog Flow dirapikan.', 'Audit front end (V3, V4, V7, V8, V9).', 'Tampilan saja.'),
    ('2026-10-31', 'Semua dialog indikator', 'Kontrol yang dipilih memakai garis abu-abu tipis dan cincin samar, bukan '
     'garis hitam 2 px.', 'User: "selected state bisa dirubah bordernya tidak setebal sekarang. Dibuat soft saja".',
     'Tampilan saja.'),
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def formulas():
    """(conventions, rows) from FORMULAS.md: the bullet text of sections 1-2 and the table of section 3."""
    text = FORMULAS.read_text()
    conv = []
    part = text.split('## 3. Formulas')[0]
    for block in re.split(r'\n(?=## )', part):
        head = block.splitlines()[0].lstrip('# ').strip()
        for item in re.split(r'\n(?=- )', block):
            if item.startswith('- '):
                conv.append((head, ' '.join(line.strip() for line in item[2:].splitlines())))
    rows = [[c.strip() for c in line.strip().strip('|').split(' | ')]
            for line in text.split('## 3. Formulas')[1].splitlines() if re.match(r'^\| FL\d\d ', line)]
    return conv, rows


def style_head(ws, widths):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='171717')
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'


def main(out):
    check = subprocess.run([sys.executable, str(ROOT / 'scripts/build_flow_registry.py'), '--check'],
                           capture_output=True, text=True)
    if check.returncode:
        raise SystemExit('the registry differs from the spreadsheet: ' + (check.stdout + check.stderr)[-500:])
    conv, rows = formulas()
    if len(rows) != 39:
        raise SystemExit(f'FORMULAS.md has {len(rows)} indicator rows, expected FL01-FL39')
    wb = openpyxl.load_workbook(SPEC)

    ws = wb.create_sheet('Rumus_Kode', 0)
    ws.append(['Id', 'Nama', 'Output', 'Rumus (seperti dihitung di Terminal)'])
    for r in rows:
        ws.append([v.replace('`', '') for v in r])
    ws.append([])
    ws.append(['Bagian', 'Konvensi bersama', '', ''])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    for head, item in conv:
        ws.append([head, item.replace('`', '').replace('**', '')])
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=4)
    style_head(ws, [8, 34, 16, 110])

    ws = wb.create_sheet('Perubahan', 1)
    ws.append(['Tanggal', 'Indikator', 'Perubahan', 'Asal', 'Pengaruh ke angka'])
    for c in CHANGES:
        ws.append(list(c))
    style_head(ws, [12, 22, 60, 40, 50])

    ws = wb.create_sheet('Versi', 2)
    ws.append(['Hal', 'Nilai'])
    for k, v in [
        ('Dibuat', date.today().isoformat()),
        ('Kontrak (sheet asli, tidak diubah)', f'{SPEC.name} sha256 {sha(SPEC)}'),
        ('Rumus', f'{FORMULAS.relative_to(ROOT)} sha256 {sha(FORMULAS)}'),
        ('Kode hitung', f'{FLOW_TS.relative_to(ROOT)} sha256 {sha(FLOW_TS)}'),
        ('Registry Terminal', f'{REGISTRY.relative_to(ROOT)} sha256 {sha(REGISTRY)}'),
        ('Cek registry = spreadsheet', (check.stdout or 'ok').strip()),
        ('Isi', 'Rumus_Kode, Perubahan dan Versi ditambahkan; sheet lain adalah spreadsheet yang disepakati 2026-10-30.'),
    ]:
        ws.append([k, v])
    style_head(ws, [34, 110])
    wb.save(out)
    print(f'{out}: {len(rows)} indicators, {len(conv)} conventions, {len(CHANGES)} changes')


if __name__ == '__main__':
    main(sys.argv[1])
