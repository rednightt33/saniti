import openpyxl, sys
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb['02_PARAMETERS']
rows = list(ws.iter_rows(values_only=True))
for r in rows[1:]:
    if r[0] and r[0] <= 'I27':
        print(r[0], r[1], r[2], 'def=', r[6], r[7], 'min=', r[8], 'max=', r[9], 'opt=', r[11], '|', (r[14] or '')[:120])
ws = wb['01_INDICATORS']
for r in list(ws.iter_rows(values_only=True))[1:]:
    if r[0] and r[0] <= 'I27':
        print(r[0], r[3], 'readiness=', r[9], 'note=', (r[10] or '')[:100])
