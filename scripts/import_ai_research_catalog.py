#!/usr/bin/env python3
"""Load the reviewed private workbook into the new research catalog, atomically.

Run only after the research-catalog schema migration and a live database preflight.
Requires openpyxl and psycopg; DATABASE_URL must identify the inspected dev database.
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
from pathlib import Path

import openpyxl


EXPECTED_SHA256 = "49b1bfff36674d6442b984e2321b2d5ad3c81e288204a119ede220aa84a6c1e1"
TABLE = 'public."AI_research_catalog"'
FIELDS = (
    "method_id", "method_name", "category", "purpose", "example_question", "analysis_kind",
    "input_grain", "required_inputs_json", "optional_inputs_json", "future_outcome_required",
    "supports_numeric_directly", "preprocessing", "parameter_keys_json", "expected_outputs_json",
    "validation_requirements_json", "main_risks", "compute_strategy", "tool_or_library_examples",
    "implementation_status",
)
JSON_FIELDS = {field for field in FIELDS if field.endswith("_json")}
BOOL_FIELDS = {"future_outcome_required", "supports_numeric_directly"}


def reviewed_rows(path: Path) -> list[dict[str, object]]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError("Workbook SHA256 differs from the reviewed attachment")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["AI_Research_Catalog"]
        records = sheet.values
        if tuple(next(records)) != FIELDS:
            raise ValueError("Workbook column names/order differ from the reviewed schema")
        result: list[dict[str, object]] = []
        for values in records:
            if not any(value is not None for value in values):
                continue
            if len(values) != len(FIELDS) or any(value is None for value in values):
                raise ValueError("Workbook contains a missing field or extra column")
            row = dict(zip(FIELDS, values))
            for field in JSON_FIELDS:
                value = json.loads(row[field])
                if not isinstance(value, list):
                    raise ValueError(f"Expected a JSON array in {field}")
                row[field] = value
            for field in BOOL_FIELDS:
                if not isinstance(row[field], bool):
                    raise ValueError(f"Expected a boolean in {field}")
            if row["implementation_status"] != "REFERENCE_ONLY":
                raise ValueError("The reviewed catalog contains a non-reference method")
            result.append(row)
        if len(result) != 18 or len({r["method_id"] for r in result}) != 18:
            raise ValueError("Expected exactly 18 distinct research methods")
        return result
    finally:
        book.close()


def import_rows(dsn: str, rows: list[dict[str, object]]) -> None:
    import psycopg
    from psycopg.types.json import Jsonb

    quoted = ", ".join(f'"{field}"' for field in FIELDS)
    placeholders = ", ".join("%s" for _ in FIELDS)
    with psycopg.connect(dsn) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SET LOCAL lock_timeout = '5s'")
            cursor.execute("SET LOCAL statement_timeout = '1min'")
            cursor.execute(f"LOCK TABLE {TABLE} IN EXCLUSIVE MODE")
            columns = cursor.execute("""
                SELECT column_name FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = 'AI_research_catalog'
                ORDER BY ordinal_position
            """).fetchall()
            if tuple(name for (name,) in columns) != FIELDS:
                raise ValueError("Live table columns differ from the reviewed workbook")
            if cursor.execute(f"SELECT count(*) FROM {TABLE}").fetchone()[0] != 0:
                raise ValueError("Research catalog is not empty; inspect live rows before import")
            for row in rows:
                values = tuple(Jsonb(row[field]) if field in JSON_FIELDS else row[field] for field in FIELDS)
                cursor.execute(f"INSERT INTO {TABLE} ({quoted}) VALUES ({placeholders})", values)
            actual = cursor.execute(f"SELECT {quoted} FROM {TABLE} ORDER BY method_id").fetchall()
            expected = [tuple(row[field] for field in FIELDS)
                        for row in sorted(rows, key=lambda entry: entry["method_id"])]
            if actual != expected:
                raise ValueError("Database readback differs from reviewed workbook; transaction rolled back")
    print(f"Imported and verified {len(rows)} rows / {len(FIELDS)} columns; source SHA256 {EXPECTED_SHA256}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: DATABASE_URL=... python scripts/import_ai_research_catalog.py WORKBOOK.xlsx")
    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        raise SystemExit("DATABASE_URL is required")
    import_rows(dsn, reviewed_rows(Path(sys.argv[1])))


if __name__ == "__main__":
    main()
