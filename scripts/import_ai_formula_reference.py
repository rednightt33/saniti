#!/usr/bin/env python3
"""Load the reviewed private workbook into the new formula reference catalog, atomically.

Run only after the formula-reference schema migration and a live database preflight.
Requires openpyxl and psycopg; DATABASE_URL must identify the inspected dev database.
"""

from __future__ import annotations

import hashlib
import os
import sys
from pathlib import Path

import openpyxl


EXPECTED_SHA256 = "09326f986728a97b8e5c885c46c8270634b541d975f7090067f184ae3b22a30d"
TABLE = 'public."AI_formula_reference"'
SOURCE_FIELDS = (
    "calculation_id", "calculation_name", "description", "required_inputs",
    "formula / method", "parameters", "output", "implementation",
)
DB_FIELDS = (
    "calculation_id", "calculation_name", "description", "required_inputs",
    "formula_method", "parameters", "output", "implementation",
)


def reviewed_rows(path: Path) -> list[dict[str, object]]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError("Workbook SHA256 differs from the reviewed attachment")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Calculation_Catalog"]
        records = sheet.values
        if tuple(next(records)) != SOURCE_FIELDS:
            raise ValueError("Workbook column names/order differ from the reviewed schema")
        result: list[dict[str, object]] = []
        seen_ids: set[str] = set()
        for values in records:
            if not any(value is not None for value in values):
                continue
            if len(values) != len(SOURCE_FIELDS) or any(
                value is None or (isinstance(value, str) and not value.strip()) for value in values
            ):
                raise ValueError("Workbook contains a missing field or extra column")
            row = dict(zip(DB_FIELDS, values))
            if row["calculation_id"] in seen_ids:
                raise ValueError(f"Duplicate calculation_id: {row['calculation_id']}")
            seen_ids.add(row["calculation_id"])
            result.append(row)
        if len(result) != 200:
            raise ValueError(f"Expected exactly 200 distinct formulas, found {len(result)}")
        return result
    finally:
        book.close()


def import_rows(dsn: str, rows: list[dict[str, object]]) -> None:
    import psycopg

    quoted = ", ".join(f'"{field}"' for field in DB_FIELDS)
    placeholders = ", ".join("%s" for _ in DB_FIELDS)
    with psycopg.connect(dsn) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SET LOCAL lock_timeout = '5s'")
            cursor.execute("SET LOCAL statement_timeout = '1min'")
            cursor.execute(f"LOCK TABLE {TABLE} IN EXCLUSIVE MODE")
            columns = cursor.execute("""
                SELECT column_name FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = 'AI_formula_reference'
                ORDER BY ordinal_position
            """).fetchall()
            if tuple(name for (name,) in columns) != DB_FIELDS:
                raise ValueError("Live table columns differ from the reviewed workbook")
            if cursor.execute(f"SELECT count(*) FROM {TABLE}").fetchone()[0] != 0:
                raise ValueError("Formula reference is not empty; inspect live rows before import")
            for row in rows:
                values = tuple(row[field] for field in DB_FIELDS)
                cursor.execute(f"INSERT INTO {TABLE} ({quoted}) VALUES ({placeholders})", values)
            actual = cursor.execute(f"SELECT {quoted} FROM {TABLE} ORDER BY calculation_id").fetchall()
            expected = [tuple(row[field] for field in DB_FIELDS)
                        for row in sorted(rows, key=lambda entry: entry["calculation_id"])]
            if actual != expected:
                raise ValueError("Database readback differs from reviewed workbook; transaction rolled back")
    print(f"Imported and verified {len(rows)} rows / {len(DB_FIELDS)} columns; source SHA256 {EXPECTED_SHA256}")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: DATABASE_URL=... python scripts/import_ai_formula_reference.py WORKBOOK.xlsx")
    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        raise SystemExit("DATABASE_URL is required")
    import_rows(dsn, reviewed_rows(Path(sys.argv[1])))


if __name__ == "__main__":
    main()
